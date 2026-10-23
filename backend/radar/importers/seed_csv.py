"""Import the normalized findings CSV (one row per entity, columns as in findings_seed.csv).

The workbook importer converts the original sheets into the same row format and calls
`import_rows`, so both paths share mapping, fixes and the data-quality report.
"""

from __future__ import annotations

import csv
import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Iterable

from sqlalchemy.orm import Session

from radar.importers.common import file_sha256, map_status, parse_date, parse_int, split_list
from radar.models import ImportBatch
from radar.normalize import NotAUrl, clean_raw_url, identify, is_valid_url
from radar.store import FindingIn, add_link, upsert_finding

ENTITY_TYPE_MAP = {"instagram_profil": "instagram_account", "instagram_reels": "instagram_reel"}
SHEET_TAGS = {"Kohne Tapintilar": "historical", "App store": "app_landing"}


def _is_live(v: str | None) -> bool | None:
    return {"yes": True, "no": False}.get((v or "").strip().lower())


def row_to_finding(row: dict, origin: str, seen_at: datetime | None) -> tuple[FindingIn, list[str]]:
    """Map one row to FindingIn. Returns (finding, fixes applied)."""
    fixes: list[str] = []
    raw = json.loads(row["raw_json"]) if row.get("raw_json") else {}
    url_in = (row.get("url") or "").strip()
    url = clean_raw_url(url_in)
    if not url:
        raise NotAUrl(url_in)
    if url_in.startswith(("/", "=")):
        fixes.append("url_repaired")
    elif not url_in.lower().startswith(("http://", "https://")):
        fixes.append("url_scheme_added")

    sheets = split_list(row.get("source_sheets", "").replace(" + ", ","))
    tags = split_list(row.get("tags"))
    for s in sheets:
        if s in SHEET_TAGS and SHEET_TAGS[s] not in tags:
            tags.append(SHEET_TAGS[s])
    if row.get("entity_type") == "app_landing_page" and "app_landing" not in tags:
        tags.append("app_landing")
    if row.get("licensed_brand_mentions"):
        tags.append("licensed_brand_misuse")
    tags.append("imported")  # backfill: excluded from "new findings" counts
    tags = [t for t in dict.fromkeys(tags) if t != "merged_duplicate"]

    status, live = map_status(row.get("original_status"), row.get("complaint_status"),
                              default=row.get("status") or "confirmed")
    is_live = _is_live(row.get("is_live"))
    if is_live is None:
        is_live = live

    landing_raw = raw.get("Landing") if "Landing" in raw else (row.get("landing_url") or None)
    landing_url = clean_raw_url(row.get("landing_url")) if is_valid_url(row.get("landing_url")) else None
    if landing_raw and not landing_url:
        fixes.append("landing_not_url_kept_raw")

    entity_type = ENTITY_TYPE_MAP.get(row.get("entity_type") or "", row.get("entity_type") or None)
    # the CSV's own source names are coarse (e.g. "tiktok"); identify() decides the platform source
    ref = identify(url)
    if entity_type in ("app_landing_page",) or ref.source == "organic_domain":
        entity_type = entity_type if entity_type in ("app_landing_page",) else "domain"
    if ref.source in ("telegram_channel",) and entity_type not in ("telegram_group", "telegram_bot", "telegram_channel"):
        entity_type = ref.entity_type
    if ref.source == "meta_ad":
        entity_type = "ad"
    if ref.source.startswith(("instagram_", "tiktok_", "facebook_", "play_store", "app_store")):
        entity_type = ref.entity_type  # the URL is more reliable than the sheet's "Növ" column

    manual = {k: v for k, v in {
        "hosting_registrar": row.get("hosting_registrar"),
        "cloudflare": row.get("cloudflare"),
        "abuse_email": row.get("abuse_email"),
        "language_geo": row.get("language_geo"),
    }.items() if v}

    data_fixes = split_list(row.get("data_fixes"))
    fixes.extend(data_fixes)
    review_reasons = split_list(row.get("review_reasons"))
    published = parse_date(row.get("start_date"))
    fin = FindingIn(
        url=url,
        entity_type=entity_type,
        handle=(row.get("handle") or "").lstrip("@") or None,
        title=row.get("name") or None,
        text=row.get("text") or None,
        landing_url=landing_url,
        landing_raw=landing_raw,
        advertiser=row.get("advertiser") or None,
        seen_at=published or seen_at,
        published_at=published,
        brands=split_list(row.get("brands")),
        matched_keywords=split_list(row.get("matched_keywords")),
        status=status,
        original_status=row.get("original_status") or None,
        complaint_status=row.get("complaint_status") or None,
        is_live=is_live,
        category=row.get("category") or None,
        suggested_action=row.get("suggested_action") or None,
        tags=tags,
        needs_review=(row.get("needs_review") or "").lower() == "yes",
        review_reasons=review_reasons,
        is_impersonation="brand_impersonation_domain" in tags,
        subscribers=parse_int(row.get("subscribers")),
        raw={
            "import": origin,
            "row_id": row.get("finding_id"),
            "sheets": sheets,
            "rows": row.get("source_rows"),
            "query": row.get("query") or None,
            "note": row.get("source_note") or None,
            "games": row.get("games") or None,
            "licensed_brand_mentions": row.get("licensed_brand_mentions") or None,
            "cluster_id": row.get("cluster_id") or None,
            "linked_entity": row.get("linked_entity") or None,
            "data_fixes": data_fixes or None,
            "original": raw,
        },
        origin=origin,
        public_id=row.get("finding_id") or None,
    )
    fin.raw = {k: v for k, v in fin.raw.items() if v is not None}
    if manual:
        fin.enrichment = {"manual": manual}
    return fin, fixes


def import_rows(session: Session, rows: Iterable[dict], *, kind: str, filename: str, sha256: str,
                seen_at: datetime | None = None) -> ImportBatch:
    batch = ImportBatch(kind=kind, filename=filename, sha256=sha256)
    session.add(batch)
    session.flush()
    origin = f"import:{kind}:{batch.id}"

    stats = Counter()
    fixes = Counter()
    skipped: list[dict] = []
    status_mismatch: list[dict] = []
    keys_in_file: dict[str, list[str]] = {}
    by_sheet = Counter()

    for row in rows:
        stats["rows_read"] += 1
        try:
            fin, row_fixes = row_to_finding(row, origin, seen_at)
            key = fin.resolve().key
        except (NotAUrl, ValueError) as exc:
            stats["skipped_invalid_url"] += 1
            skipped.append({"row": row.get("finding_id") or row.get("source_rows"), "url": row.get("url"), "error": str(exc)})
            continue
        fixes.update(row_fixes)
        if row.get("status") and row["status"] != fin.status:
            status_mismatch.append({"row": row.get("finding_id"), "csv": row["status"], "mapped": fin.status})
        keys_in_file.setdefault(key, []).append(row.get("finding_id") or "")
        res = upsert_finding(session, fin, status_policy="max")
        f = res.finding
        if fin.raw.get("cluster_id"):
            add_link(session, f.id, "ad_cluster", value=fin.raw["cluster_id"])
        if res.created:
            stats["created"] += 1
            for s in fin.raw.get("sheets", []):
                by_sheet[s] += 1
        else:
            stats["merged_into_existing"] += 1
        # entries merged before export (tag merged_duplicate / "A + B" source rows)
        if "+" in (row.get("source_rows") or ""):
            stats["pre_merged_rows"] += (row.get("source_rows") or "").count("+")

    dupes = {k: v for k, v in keys_in_file.items() if len(v) > 1}
    report = {
        "rows_read": stats["rows_read"],
        "created": stats["created"],
        "merged_into_existing": stats["merged_into_existing"],
        "duplicates_in_file": sum(len(v) - 1 for v in dupes.values()),
        "duplicate_groups": [{"key": k, "rows": v} for k, v in list(dupes.items())[:200]],
        "duplicates_merged_before_export": stats["pre_merged_rows"],
        "skipped_invalid_url": skipped,
        "fixes": dict(fixes),
        "status_mismatch": status_mismatch[:200],
        "created_by_sheet": dict(by_sheet),
    }
    batch.report = report
    session.flush()
    return batch


def read_seed_xlsx(path: Path) -> list[dict]:
    """The same seed as an Excel workbook: rows live on the "Findings" sheet; other sheets are summaries."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = next((wb[n] for n in wb.sheetnames if n.strip().lower() == "findings"), None)
    if ws is None:
        raise ValueError(f"{path.name}: no 'Findings' sheet")
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    out = []
    for r in rows:
        rec = {h: ("" if v is None else str(v)) for h, v in zip(header, r) if h}
        if any(rec.values()):
            out.append(rec)
    return out


def import_seed_csv(session: Session, path: Path, seen_at: datetime | None = None) -> ImportBatch:
    """Import the normalized seed from .csv or .xlsx (sheet "Findings")."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        rows, kind = read_seed_xlsx(path), "seed_xlsx"
    else:
        with open(path, encoding="utf-8-sig", newline="") as fh:
            rows, kind = list(csv.DictReader(fh)), "seed_csv"
    return import_rows(session, rows, kind=kind, filename=path.name, sha256=file_sha256(path), seen_at=seen_at)
