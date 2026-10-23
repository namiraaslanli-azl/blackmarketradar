"""Original workbook import, including the known data issues listed in the brief."""

from openpyxl import Workbook
from sqlalchemy import select

from radar.importers.workbook import import_workbook, telegram_url
from radar.models import Finding


def make_workbook(path):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Organic Domain")
    ws.append(["Qara Bazar Monitorinq Hesabatı"])  # title row above the header
    ws.append(["Domen / URL", "Kateqoriya", "Hosting / Registrator", "Cloudflare?", "Abuse Email",
               "Təklif Olunan Təsir Mexanizmi"])
    ws.append(["aze.1xbet.com", "Birbaşa Operativ Mirror", "Cloudflare Inc.", "Bəli", "abuse@cloudflare.com",
               "İKTA DNS/IP Block + Cloudflare Abuse"])
    ws.append(["https://azmostbet90.com/"])
    ws = wb.create_sheet("Kohne Tapintilar")
    ws.append(["Saytlar", "Status"])
    ws.append(["/https://www.aze.1xbet.com/", "Passiv"])
    ws.append(["https://azerlotereya.org/", "Aktiv"])
    ws.append(["AZMOSTBET90.com", "Bloklanıb"])
    ws = wb.create_sheet("App store")
    ws.append(["Link"])
    ws.append(["Fragaplay.com/app"])
    ws = wb.create_sheet("Meta Ads")
    ws.append(["Reklamçı", "Reklam linki", "Başlama tarixi", "Landing", "Mətn (qısa)", "Niyə bu sinif",
               "Şikayət statusu", "Status"])
    ws.append(["Lake Smith Farm", "https://www.facebook.com/ads/library/?id=111", "27.10.2024",
               "http://play.google.com/store/apps/details?id=com.unearby.sayh",
               "Mega qeydiyyat bonusu 125% depozit + 250 pulsuz fırlanma", "", "Şikayət olunub", "Bloklanıb"])
    ws.append(["Ojobob online shop", "https://www.facebook.com/ads/library/?id=222", "28.10.2024",
               "Facebook, Instagram", "Kazino bonus", "", "Şikayət olunub", "Aktiv"])
    ws.append(["X", "https://www.facebook.com/ads/library/?id=333", "", "", "", "", "", "Aktiv"])
    ws = wb.create_sheet("Telegram")
    ws.append(["İstifadəçi adı", "Ad", "Link", "Növ", "Abunəçi / üzv", "Dil / geo", "Əsas məlumat və əlaqələr",
               "Mənbə", "Status"])
    ws.append(["=+TgTtogD6uk8xMjQ0", "Söhbət qrupu | PINCO", "", "Qrup (dəvət linki)", "1677", "AZ",
               "pin-co.cc/vMwTt, əlavə kanal +v4bKeObN8DdhMWRk", "", "Yoxlanılmayıb"])
    ws.append(["melbetkassa", "Kassa", "", "Kanal", "1.2k", "AZ", "Kassa: M10, Kapital Bank, USDT", "", "Aktiv"])
    ws = wb.create_sheet("Tiktok Organic")
    ws.append(["User / Channel", "URL"])
    ws.append([".1winaze (1win Gelin)", "https://www.tiktok.com/@.1winaze"])
    wb.save(path)


def test_workbook_import(session, tmp_path):
    p = tmp_path / "Qara_Bazar_Monitorinq_Hesabatı_-_Weekly.xlsx"
    make_workbook(p)
    batch = import_workbook(session, p)
    rep = batch.report
    assert rep["duplicates_in_file"] == 2  # aze.1xbet.com and azmostbet90.com across sheets

    by_key = {f.canonical_key: f for f in session.scalars(select(Finding))}
    xbet = by_key["domain:aze.1xbet.com"]
    assert xbet.category == "Birbaşa Operativ Mirror"
    assert xbet.suggested_action.startswith("İKTA")
    assert xbet.enrichment["manual"]["abuse_email"] == "abuse@cloudflare.com"
    assert xbet.original_status == "Passiv" and xbet.is_live is False
    assert by_key["domain:azmostbet90.com"].status == "blocked"

    app = by_key["domain:fragaplay.com"]
    assert app.source == "organic_domain" and "app_landing" in app.tags

    ad1, ad2, ad3 = by_key["metaad:111"], by_key["metaad:222"], by_key["metaad:333"]
    assert ad1.landing_url.startswith("http://play.google.com/") and ad1.status == "blocked"
    assert ad1.complaint_status == "Şikayət olunub"
    assert ad2.landing_url is None and ad2.landing_raw == "Facebook, Instagram"
    assert ad2.status == "reported"
    assert ad3.landing_url is None and ad3.landing_raw is None

    tg = by_key["tg:+TgTtogD6uk8xMjQ0"]
    assert tg.url == "https://t.me/+TgTtogD6uk8xMjQ0"
    assert tg.entity_type == "telegram_group" and tg.subscribers == 1677
    assert tg.needs_review
    assert by_key["tg:melbetkassa"].subscribers == 1200
    assert by_key["ttacc:.1winaze"].title == "1win Gelin"
    assert rep["fixes"]["fixed_excel_formula_username"] == 1


def test_telegram_url_rebuild():
    assert telegram_url("=+QL1-qIWYVMQwZjY0", "") == ("https://t.me/+QL1-qIWYVMQwZjY0", "fixed_excel_formula_username")
    assert telegram_url("Betpriz_bot", "https://t.me/Betpriz_bot") == ("https://t.me/Betpriz_bot", None)
    assert telegram_url("@aze1xbetkassa", "") == ("https://t.me/aze1xbetkassa", "link_rebuilt_from_username")
